"""Offline regression checks in an isolated directory under artifacts/; no OPS access."""
import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
(ROOT / 'artifacts/validation').mkdir(parents=True, exist_ok=True)
with tempfile.TemporaryDirectory(dir=ROOT / 'artifacts/validation') as tmp:
    root = Path(tmp)
    shutil.copytree(ROOT / 'scripts', root / 'scripts')
    shutil.copytree(ROOT / 'config', root / 'config', ignore=shutil.ignore_patterns('env.md'))
    (root / 'scenario/example').mkdir(parents=True)
    (root / 'results').mkdir()
    scenarios = [{'id': f'SCN-{i:04}', 'title': 'Synthetic framework validation', 'category': 'positive',
                  'priority': 'low', 'screen': '', 'requirements': []} for i in range(1, 6)]
    (root / 'scenario/example/example_scenarios.json').write_text(json.dumps({'scenarios': scenarios}))
    specs = []
    for s, status, error in zip(scenarios, ['passed','failed','failed','failed','skipped'], ['', '[AUTH_NOT_VALIDATED] pending', '[AUTH_FAILURE] failed', 'Assertion failed', '']):
        specs.append({'title': s['id'] + ': ' + s['title'], 'tests': [{'results': [{'status': status, 'errors': [{'message': error}] if error else []}]}]})
    report = {'stats': {'startTime': '2026-09-23T00:00:00Z', 'duration': 1}, 'suites': [{'specs': specs}]}
    (root / 'input.json').write_text(json.dumps(report))
    def run(*args):
        return subprocess.run(args, cwd=root, check=True, capture_output=True, text=True).stdout
    out = run('python3','scripts/playwright_to_results.py','input.json','example').strip()
    data = json.loads((root / out).read_text())
    assert [s['status'] for s in data['scenarios']] == ['passed','blocked','blocked','failed','skipped']
    assert data['environment']['user'] is None
    assert data['scenarios'][0]['title'] == scenarios[0]['title']
    run('python3','scripts/generate_report.py',out)
    wb = load_workbook(next((root / 'reports').glob('*.xlsx')))
    assert wb.sheetnames == ['Summary', 'Detail', 'Failed_BugCandidates']
    assert wb['Detail'].max_row == 6
    assert wb['Summary']['A1'].value.endswith('OPS')
    for i in [1,2]:
        (root / f'results/_batch0{i}__example__check.json').write_text(json.dumps({'scenarios': [data['scenarios'][0]]}))
    run('python3','scripts/merge_batches.py','example','check')
    merged = json.loads((root / 'results/example__check.json').read_text())
    assert len(merged['scenarios']) == 5  # duplicate deduplication + missing scenarios
    assert merged['environment'] == {'baseUrl':'', 'user':'', 'role':''}
    run('python3','scripts/generate_report.py','results/example__check.json')
print('PASS: converter auth classification, metadata, empty credentials, Excel, merge schema/duplicates/missing/environment')
